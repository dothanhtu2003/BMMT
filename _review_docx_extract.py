from docx import Document
from pathlib import Path
import sys

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

p = Path(r"C:\Users\Administrator\Downloads\BMTT_N01_G13_52200248_52200286_52200240 (1).docx")
d = Document(p)

print("PARAGRAPHS", len(d.paragraphs))
print("TABLES", len(d.tables))
rel_imgs = [r for r in d.part.rels.values() if "image" in r.reltype]
print("IMAGES", len(rel_imgs))
print("---TEXT---")
for i, para in enumerate(d.paragraphs, 1):
    t = para.text.strip()
    if t:
        print(f"{i:04d}: [{para.style.name}] {t}")

print("---TABLES---")
for ti, table in enumerate(d.tables, 1):
    print(f"TABLE {ti} rows={len(table.rows)} cols={len(table.columns)}")
    for row in table.rows[:10]:
        print(" | ".join(c.text.strip().replace("\n", " / ") for c in row.cells))
